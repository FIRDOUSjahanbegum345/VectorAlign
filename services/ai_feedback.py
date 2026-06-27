import os
import json
import re
import pdfplumber
import pypdf
import docx

from dotenv import load_dotenv
from google import genai

load_dotenv()

client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

# ===========================
# Resume Text Extraction
# ===========================

def extract_resume_text(file_path, file_extension):
    """
    Extract text from uploaded resume.
    Supports PDF, DOCX and TXT.
    """
    file_extension = file_extension.lower()

    if file_extension == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_extension == "docx":
        return extract_text_from_docx(file_path)
    elif file_extension == "txt":
        return extract_text_from_txt(file_path)
    else:
        raise ValueError("Unsupported file format")


def extract_text_from_pdf(file_path):
    """
    Extract text using pdfplumber.
    Falls back to pypdf if pdfplumber fails.
    """
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception:
        try:
            with open(file_path, "rb") as file:
                reader = pypdf.PdfReader(file)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception:
            return ""
    return text.strip()


def extract_text_from_docx(file_path):
    """
    Extract text from DOCX.
    """
    document = docx.Document(file_path)
    text = []
    for para in document.paragraphs:
        if para.text.strip():
            text.append(para.text)
    return "\n".join(text)


def extract_text_from_txt(file_path):
    """
    Extract text from TXT.
    """
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()

# ===========================
# JSON Utilities
# ===========================

def extract_json(text):
    """
    Gemini sometimes wraps JSON inside ```json blocks.
    Remove them and parse safely.
    """
    if not text:
        return None

    text = text.strip()
    text = re.sub(r"^```json", "", text)
    text = re.sub(r"^```", "", text)
    text = re.sub(r"```$", "", text)
    text = text.strip()

    try:
        return json.loads(text)
    except Exception:
        return None

# ===========================
# Fallback Response
# ===========================

def fallback(ats_score):
    """
    Default response if Gemini fails.
    """
    return {
        "ats_score_evaluation": f"ATS Score: {ats_score}/100. Resume requires further optimization.",
        "missing_keywords": [],
        "structural_feedback": [
            "Improve formatting",
            "Use ATS-friendly headings",
            "Add measurable achievements",
            "Include more job-specific keywords"
        ],
        "quantifiable_impact_suggestions": [
            "Mention percentage improvements",
            "Highlight measurable project outcomes"
        ],
        "tailored_summary_suggestion":
        "AI Engineer passionate about software development, machine learning and scalable applications.",
        "optimized_resume": """
PROFESSIONAL SUMMARY
• Motivated Software Engineer with experience in Python, JavaScript, HTML, CSS and Machine Learning.
• Strong knowledge of FastAPI, AI, Data Science and Web Development.
• Experienced in developing scalable applications and AI-powered solutions.
• Passionate about solving real-world engineering problems.

TECHNICAL SKILLS
• Python
• Java
• C
• HTML5
• CSS3
• JavaScript
• FastAPI
• SQL
• Machine Learning
• Pandas
• NumPy
• Scikit-learn
• Git
• GitHub

WORK EXPERIENCE
• Add your professional experience here.

PROJECTS
• AI Resume Analyzer
• Network Security Scanner

EDUCATION
Bachelor's Degree
Artificial Intelligence & Data Science

CERTIFICATIONS
• Add certifications here

CORE STRENGTHS
• Problem Solving
• Teamwork
• Communication
• Leadership
• Critical Thinking
""",
        "keyword_match_score": ats_score,
        "job_match_score": ats_score
    }

# ===========================
# AI Feedback Generator
# ===========================

def generate_feedback(resume_text, job_description, ats_score):
    prompt = f"""
You are a Senior ATS Resume Writer, Resume Reviewer and Technical Recruiter.

Analyze the resume against the job description.

Return ONLY valid JSON.

The JSON MUST follow this EXACT schema:

{{
    "ats_score_evaluation":"",
    "missing_keywords":[],
    "structural_feedback":[],
    "quantifiable_impact_suggestions":[],
    "tailored_summary_suggestion":"",
    "optimized_resume":"",
    "keyword_match_score":0,
    "job_match_score":0
}}

STRICT RULES:
1. Return ONLY JSON.
2. Do NOT use markdown.
3. Do NOT wrap JSON inside ``` blocks.
4. Do NOT explain anything.
5. Do NOT include extra text.
6. optimized_resume must be a COMPLETE ATS-friendly resume.
7. Resume should be around 700-1200 words.
8. Naturally include missing keywords.
9. Improve grammar, formatting and readability.
10. Use professional recruiter language.

optimized_resume MUST contain these sections:

PROFESSIONAL SUMMARY
• 4-6 bullet points

TECHNICAL SKILLS
• Bullet list

WORK EXPERIENCE
• Bullet points
• Strong action verbs
• Quantifiable achievements wherever possible

PROJECTS
• Mention technologies used
• Mention business impact
• Bullet points

EDUCATION

CERTIFICATIONS
• If unavailable generate suitable placeholders.

CORE STRENGTHS
• Bullet points

ATS KEYWORDS
• Important keywords from the Job Description.

ATS SCORE:
{ats_score}

RESUME:
{resume_text}

JOB DESCRIPTION:
{job_description}
"""

    models = [
        "gemini-2.5-flash",
        "gemini-2.5-flash-lite",
        "gemini-2.0-flash"
    ]

    for model in models:
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt
            )
            text = response.text
            data = extract_json(text)

            if data:
                if "keyword_match_score" not in data:
                    data["keyword_match_score"] = ats_score
                if "job_match_score" not in data:
                    data["job_match_score"] = ats_score
                if "missing_keywords" not in data:
                    data["missing_keywords"] = []
                if "structural_feedback" not in data:
                    data["structural_feedback"] = []
                if "quantifiable_impact_suggestions" not in data:
                    data["quantifiable_impact_suggestions"] = []
                if "tailored_summary_suggestion" not in data:
                    data["tailored_summary_suggestion"] = ""
                if "optimized_resume" not in data:
                    data["optimized_resume"] = ""
                if "ats_score_evaluation" not in data:
                    data["ats_score_evaluation"] = f"ATS Score: {ats_score}/100."

                return data

        except Exception as e:
            print(f"Gemini Error ({model}): {e}")

    return fallback(ats_score)
